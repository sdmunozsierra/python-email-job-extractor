from docx import Document
from docx.shared import Inches

# Create a new Word document
doc = Document()

# Add a table to the document
table = doc.add_table(rows=4, cols=6)  # Creating a table with 4 rows and 6 columns

# Set the style of the table for better visibility
table.style = 'Table Grid'

# Define the content for each cell
content = [
    ["SUMM1", "", "SUMM3", "", "SUMM5", ""],  # First row
    ["POS1", "", "POS3", "", "POS5", ""],      # Second row
    ["", "POS2", "", "POS4", "", ""],          # Third row
    ["", "SUMM2", "", "SUMM4", "", ""]         # Fourth row
]

# Populate the table with the content
for row_num, row_content in enumerate(content):
    for col_num, cell_content in enumerate(row_content):
        table.cell(row_num, col_num).text = cell_content


# Save the document
doc.save('professional_timeline_cv.docx')
