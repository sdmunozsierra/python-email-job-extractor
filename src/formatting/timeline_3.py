from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# Initialize a new Word document
doc = Document()

# Sample data for the timeline
experiences = [
    {"title": "JOB1", "summary": "SUM1"},
    {"title": "JOB2", "summary": "SUM2"},
    {"title": "JOB3", "summary": "SUM3"},
    # ... add more experiences as needed
]

# Calculate the number of columns (2 per job experience, plus 1 for padding)
num_cols = 2 * len(experiences) + 1

# Create the table for the timeline
table = doc.add_table(rows=5, cols=num_cols)
table.style = 'Table Grid'

# Function to fill the table with job experiences
def fill_timeline_table(table, experiences):
    for i, experience in enumerate(experiences):
        col_index = 2 * i

        # Set summaries and job titles in the correct rows and columns
        if i % 2 == 0:  # Top row for odd-numbered jobs
            table.cell(0, col_index).text = experience['summary']
            table.cell(1, col_index).text = experience['title']
        else:  # Bottom row for even-numbered jobs
            table.cell(3, col_index).text = experience['title']
            table.cell(4, col_index).text = experience['summary']

# Fill the table with the experiences
fill_timeline_table(table, experiences)

# Save the document
doc.save('experience_timeline.docx')